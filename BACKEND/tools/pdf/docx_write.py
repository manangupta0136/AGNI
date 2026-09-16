from __future__ import annotations

import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional, Tuple

import docx
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

DATA_DIR = Path(__file__).resolve().parent.parent.parent / "data"
REPORTS_DIR = DATA_DIR / "reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

# ── Corporate Styling Constants ──────────────────────────────────────────────
COLOR_PRIMARY_NAVY = RGBColor(30, 58, 138)     # #1E3A8A (Headings & Table Headers)
COLOR_SECONDARY_BLUE = RGBColor(37, 99, 235)   # #2563EB (Subheadings)
COLOR_DARK_SLATE = RGBColor(51, 65, 85)        # #334155 (H3 & Subtitles)
COLOR_MUTED_SLATE = RGBColor(100, 116, 139)    # #64748B (Metadata & Captions)
COLOR_BODY_TEXT = RGBColor(31, 41, 55)         # #1F2937 (Body Text)
COLOR_WHITE = RGBColor(255, 255, 255)          # #FFFFFF (Header text)

HEX_HEADER_BG = "1E3A8A"                       # Navy
HEX_HEADER_ACCENT = "2563EB"                   # Blue accent stripe under the title banner
COLOR_LIGHT_BLUE = RGBColor(199, 216, 250)     # Banner subtext on navy
HEX_ALT_ROW_BG = "F8FAFC"                      # Light Slate Tint
HEX_CALLOUT_BG = "F1F5F9"                      # Light Gray Callout Tint
HEX_BORDER_COLOR = "CBD5E1"                    # Slate Border

FONT_FAMILY = "Calibri"
FONT_CODE = "Consolas"


def _set_cell_shading(cell, hex_color: str):
    """Set background color for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag.endswith("shd"):
            tc_pr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set cell padding in twentieths of a point (dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def _set_cell_borders(cell, color=HEX_BORDER_COLOR, sz="4"):
    """Set subtle border around a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)


def _add_styled_run(
    paragraph,
    text: str,
    bold: bool = False,
    italic: bool = False,
    is_code: bool = False,
    color: Optional[RGBColor] = None,
    size: Optional[Pt] = None,
):
    """Add a run with specific font properties to a paragraph."""
    run = paragraph.add_run(text)
    run.font.name = FONT_CODE if is_code else FONT_FAMILY
    run.font.size = size or (Pt(9.5) if is_code else Pt(10.5))
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    elif is_code:
        run.font.color.rgb = RGBColor(15, 23, 42)
    else:
        run.font.color.rgb = COLOR_BODY_TEXT
    return run


def _append_formatted_inline_text(paragraph, text: str, default_color: Optional[RGBColor] = None, default_size: Optional[Pt] = None):
    """
    Parse inline markdown:
    - ***bold italic***
    - **bold**
    - *italic*
    - `inline code`
    """
    # Regex tokenizes by code, bold-italic, bold, italic
    token_pattern = re.compile(
        r"(`[^`]+`|\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|___[^_]+___|__[^_]+__|_[^_]+_)"
    )
    parts = token_pattern.split(text)

    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            _add_styled_run(paragraph, part[1:-1], is_code=True)
        elif (part.startswith("***") and part.endswith("***") and len(part) >= 6) or \
             (part.startswith("___") and part.endswith("___") and len(part) >= 6):
            _add_styled_run(paragraph, part[3:-3], bold=True, italic=True, color=default_color, size=default_size)
        elif (part.startswith("**") and part.endswith("**") and len(part) >= 4) or \
             (part.startswith("__") and part.endswith("__") and len(part) >= 4):
            _add_styled_run(paragraph, part[2:-2], bold=True, color=default_color, size=default_size)
        elif (part.startswith("*") and part.endswith("*") and len(part) >= 2) or \
             (part.startswith("_") and part.endswith("_") and len(part) >= 2):
            _add_styled_run(paragraph, part[1:-1], italic=True, color=default_color, size=default_size)
        else:
            _add_styled_run(paragraph, part, color=default_color, size=default_size)


def _render_table(doc: Document, table_lines: List[str]):
    """Parse and render a Markdown table into a styled Word table."""
    parsed_rows: List[List[str]] = []
    for line in table_lines:
        line_clean = line.strip()
        if not line_clean.startswith("|"):
            continue
        # Split by pipe and remove leading/trailing empty cells
        cells = [c.strip() for c in line_clean.split("|")[1:-1]]
        # Skip markdown separator row |---|---|
        if all(re.match(r"^:?-+:?$", c) for c in cells if c):
            continue
        if cells:
            parsed_rows.append(cells)

    if not parsed_rows:
        return

    num_cols = max(len(r) for r in parsed_rows)
    num_rows = len(parsed_rows)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for r_idx, row_data in enumerate(parsed_rows):
        is_header = (r_idx == 0)
        row = table.rows[r_idx]

        for c_idx in range(num_cols):
            cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            _set_cell_borders(cell, color=HEX_BORDER_COLOR)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

            if is_header:
                _set_cell_shading(cell, HEX_HEADER_BG)
                _append_formatted_inline_text(p, cell_text, default_color=COLOR_WHITE, default_size=Pt(10))
                for run in p.runs:
                    run.bold = True
            else:
                bg = HEX_ALT_ROW_BG if r_idx % 2 == 1 else "FFFFFF"
                _set_cell_shading(cell, bg)
                _append_formatted_inline_text(p, cell_text, default_size=Pt(9.5))

    # Add spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)


def _render_code_block(doc: Document, code_lines: List[str]):
    """Render a multiline code block with shaded background and monospaced font."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _set_cell_shading(cell, HEX_CALLOUT_BG)
    _set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    _set_cell_borders(cell, color=HEX_BORDER_COLOR)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15

    code_text = "\n".join(code_lines)
    run = p.add_run(code_text)
    run.font.name = FONT_CODE
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(15, 23, 42)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(6)


def _render_callout_box(doc: Document, quote_lines: List[str]):
    """Render a blockquote/callout box with accent formatting."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _set_cell_shading(cell, HEX_CALLOUT_BG)
    _set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    # Set strong left accent border (Navy) and clear other borders
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_HEADER_BG}"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    quote_text = " ".join([q.lstrip("> ").strip() for q in quote_lines if q.strip()])
    _append_formatted_inline_text(p, quote_text, default_color=COLOR_DARK_SLATE, default_size=Pt(10))
    for r in p.runs:
        r.italic = True

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(6)


def _clear_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="nil"/><w:bottom w:val="nil"/>'
        f'<w:left w:val="nil"/><w:right w:val="nil"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)


def _set_cell_bottom_accent(cell, color=HEX_HEADER_ACCENT, sz="28"):
    """A single colored bottom border — used under the title banner as a
    thin accent stripe, matching the PDF/PPTX generators' accent bar."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)


def _add_page_number_field(paragraph):
    """Inserts a live "Page X of Y" field (not static text) into a footer
    paragraph — the standard fldChar/instrText sequence python-docx has no
    high-level API for."""
    def _field(instr):
        run = OxmlElement("w:r")
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr_text = OxmlElement("w:instrText")
        instr_text.set(qn("xml:space"), "preserve")
        instr_text.text = instr
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run.append(fld_begin)
        run.append(instr_text)
        run.append(fld_sep)
        run.append(fld_end)
        return run

    r1 = paragraph.add_run()
    r1.font.name = FONT_FAMILY
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = COLOR_MUTED_SLATE
    r1._r.append(_field(" PAGE "))

    r2 = paragraph.add_run(" of ")
    r2.font.name = FONT_FAMILY
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = COLOR_MUTED_SLATE

    r3 = paragraph.add_run()
    r3.font.name = FONT_FAMILY
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = COLOR_MUTED_SLATE
    r3._r.append(_field(" NUMPAGES "))


def docx_write(
    title: str,
    content: str,
    filename: Optional[str] = None,
    subtitle: Optional[str] = None,
    organization: str = "AGNI AIR-GAPPED INTELLIGENCE WORKBENCH",
) -> str:
    """
    Generate a high-fidelity, corporate-styled Word (.docx) document from Markdown text.

    Supports:
    - Page setup with 1-inch standard margins
    - Styled title banner, corporate header, and metadata
    - Headings (H1, H2, H3, H4) with hierarchy colors and keep_with_next
    - Bullet lists, numbered lists with proper Word list styles
    - Inline formatting: bold, italic, bold-italic, inline code
    - Native formatted Word tables from Markdown tables with headers & alternate shading
    - Code blocks with monospaced background styling
    - Callout boxes / blockquotes (> quote) with accent borders
    - Dividers (---)

    Returns the absolute file path of the created .docx file.
    """
    if not filename:
        clean_title = re.sub(r"[^\w\-_]", "_", title)[:40].strip("_") or "Report"
        filename = f"AGNI_{clean_title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    if not filename.lower().endswith(".docx"):
        filename += ".docx"

    output_path = REPORTS_DIR / filename

    doc = Document()

    # 1. Page Margins Setup
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
    section = doc.sections[0]

    # 2. Header / Title Block — a full-width colored banner (built from a
    # single-cell, border-free table, the same shading trick used for the
    # callout/code blocks below) instead of plain text at the top of a white
    # page, so the document reads as designed rather than a text dump. Same
    # navy/blue palette as pdf_generate.py and pptx_generate.py.
    content_width = section.page_width - section.left_margin - section.right_margin
    banner = doc.add_table(rows=1, cols=1)
    banner.autofit = False
    banner.columns[0].width = content_width
    banner_cell = banner.rows[0].cells[0]
    banner_cell.width = content_width
    _set_cell_shading(banner_cell, HEX_HEADER_BG)
    _set_cell_margins(banner_cell, top=200, bottom=220, left=200, right=200)
    _clear_cell_borders(banner_cell)
    _set_cell_bottom_accent(banner_cell)

    banner_p = banner_cell.paragraphs[0]
    if organization:
        org_run = banner_p.add_run(organization.upper())
        org_run.font.name = FONT_FAMILY
        org_run.font.size = Pt(8.5)
        org_run.font.bold = True
        org_run.font.color.rgb = COLOR_LIGHT_BLUE

    title_p = banner_cell.add_paragraph()
    title_p.paragraph_format.space_before = Pt(4)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run(title)
    title_run.font.name = FONT_FAMILY
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_WHITE

    meta_p = banner_cell.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(0)
    ts_str = datetime.now(timezone.utc).strftime("%B %d, %Y - %H:%M UTC")
    meta_text = f"Generated: {ts_str}  |  Confidential On-Premise Document"
    if subtitle:
        meta_text = f"{subtitle}  |  {meta_text}"
    meta_run = meta_p.add_run(meta_text)
    meta_run.font.name = FONT_FAMILY
    meta_run.font.size = Pt(9.5)
    meta_run.font.italic = True
    meta_run.font.color.rgb = COLOR_LIGHT_BLUE

    spacer_p = doc.add_paragraph()
    spacer_p.paragraph_format.space_before = Pt(0)
    spacer_p.paragraph_format.space_after = Pt(14)

    # 2b. Footer — page-numbered, present on every page, so a multi-page
    # report doesn't just trail off with no sense of length/position.
    footer_p = section.footer.paragraphs[0]
    footer_p.paragraph_format.tab_stops.add_tab_stop(content_width, WD_TAB_ALIGNMENT.RIGHT)
    footer_left = footer_p.add_run(f"{organization.upper()} — Confidential On-Premise Document\tPage ")
    footer_left.font.name = FONT_FAMILY
    footer_left.font.size = Pt(8.5)
    footer_left.font.color.rgb = COLOR_MUTED_SLATE
    _add_page_number_field(footer_p)

    # 3. Parse Markdown Body
    lines = content.splitlines()
    idx = 0
    total_lines = len(lines)

    while idx < total_lines:
        line = lines[idx]
        stripped = line.strip()

        if not stripped:
            idx += 1
            continue

        # A. Code Blocks ```
        if stripped.startswith("```"):
            code_lines = []
            idx += 1
            while idx < total_lines and not lines[idx].strip().startswith("```"):
                code_lines.append(lines[idx])
                idx += 1
            idx += 1  # Skip closing ```
            _render_code_block(doc, code_lines)
            continue

        # B. Markdown Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while idx < total_lines and lines[idx].strip().startswith("|"):
                table_lines.append(lines[idx])
                idx += 1
            _render_table(doc, table_lines)
            continue

        # C. Blockquotes / Callouts
        if stripped.startswith(">"):
            quote_lines = []
            while idx < total_lines and lines[idx].strip().startswith(">"):
                quote_lines.append(lines[idx])
                idx += 1
            _render_callout_box(doc, quote_lines)
            continue

        # D. Horizontal Divider
        if stripped in ("---", "***", "___") or re.match(r"^[-*_]{3,}$", stripped):
            div_p = doc.add_paragraph()
            div_p.paragraph_format.space_before = Pt(6)
            div_p.paragraph_format.space_after = Pt(8)
            div_run = div_p.add_run("―" * 48)
            div_run.font.size = Pt(8)
            div_run.font.color.rgb = RGBColor(226, 232, 240)
            idx += 1
            continue

        # E. Headings
        if stripped.startswith("#### "):
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(8)
            h_p.paragraph_format.space_after = Pt(2)
            h_p.paragraph_format.keep_with_next = True
            _append_formatted_inline_text(h_p, stripped[5:], default_color=COLOR_MUTED_SLATE, default_size=Pt(11))
            for r in h_p.runs:
                r.bold = True
            idx += 1
            continue

        if stripped.startswith("### "):
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(10)
            h_p.paragraph_format.space_after = Pt(3)
            h_p.paragraph_format.keep_with_next = True
            _append_formatted_inline_text(h_p, stripped[4:], default_color=COLOR_DARK_SLATE, default_size=Pt(12))
            for r in h_p.runs:
                r.bold = True
            idx += 1
            continue

        if stripped.startswith("## "):
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(14)
            h_p.paragraph_format.space_after = Pt(4)
            h_p.paragraph_format.keep_with_next = True
            _append_formatted_inline_text(h_p, stripped[3:], default_color=COLOR_SECONDARY_BLUE, default_size=Pt(14))
            for r in h_p.runs:
                r.bold = True
            idx += 1
            continue

        if stripped.startswith("# "):
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(18)
            h_p.paragraph_format.space_after = Pt(6)
            h_p.paragraph_format.keep_with_next = True
            _append_formatted_inline_text(h_p, stripped[2:], default_color=COLOR_PRIMARY_NAVY, default_size=Pt(16))
            for r in h_p.runs:
                r.bold = True
            idx += 1
            continue

        # F. Bullet Lists
        bullet_match = re.match(r"^(\s*)[-*+•]\s+(.*)$", line)
        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            bullet_text = bullet_match.group(2)
            b_p = doc.add_paragraph(style="List Bullet")
            b_p.paragraph_format.space_before = Pt(1)
            b_p.paragraph_format.space_after = Pt(2)
            if indent_spaces >= 2:
                b_p.paragraph_format.left_indent = Inches(0.25 * (indent_spaces // 2 + 1))
            _append_formatted_inline_text(b_p, bullet_text, default_size=Pt(10.5))
            idx += 1
            continue

        # G. Numbered Lists
        num_match = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if num_match:
            num_text = num_match.group(2)
            n_p = doc.add_paragraph(style="List Number")
            n_p.paragraph_format.space_before = Pt(1)
            n_p.paragraph_format.space_after = Pt(2)
            _append_formatted_inline_text(n_p, num_text, default_size=Pt(10.5))
            idx += 1
            continue

        # H. Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        _append_formatted_inline_text(p, stripped, default_size=Pt(10.5))
        idx += 1

    doc.save(str(output_path))
    return str(output_path)
