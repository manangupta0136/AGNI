from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Dict, List, Optional
from pptx import Presentation


def pptx_read(pptx_path: str) -> Dict[str, Any]:
    """
    Read and extract structured content from a PowerPoint (.pptx) file.

    Extracts:
    - Slide titles and subtitles
    - Text boxes and bullet points in visual reading order
    - PowerPoint tables formatted as clean Markdown tables
    - Speaker notes

    Returns a dictionary containing:
    - structured_text: Complete markdown representation of the presentation
    - slides: List of per-slide structured objects
    - total_slides: Number of slides
    - title: Presentation main title
    """
    resolved_path = Path(pptx_path).resolve()
    if not resolved_path.is_file():
        raise FileNotFoundError(f"PowerPoint file not found at: {pptx_path}")

    prs = Presentation(str(resolved_path))
    slides_data: List[Dict[str, Any]] = []
    markdown_sections: List[str] = []

    presentation_title = resolved_path.stem.replace("_", " ")

    for idx, slide in enumerate(prs.slides, 1):
        slide_title = ""
        slide_paragraphs: List[str] = []
        slide_tables: List[str] = []
        notes_text = ""

        # 1. Check title placeholder if available
        if slide.shapes.title and slide.shapes.title.text:
            slide_title = slide.shapes.title.text.strip()

        # 2. Sort shapes by visual top-to-bottom, left-to-right order
        shapes = sorted(
            slide.shapes,
            key=lambda s: (getattr(s, "top", 0) or 0, getattr(s, "left", 0) or 0)
        )

        for shape in shapes:
            # Check for native tables
            if shape.has_table:
                table = shape.table
                table_md_rows: List[str] = []
                headers: List[str] = []

                for r_idx, row in enumerate(table.rows):
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if r_idx == 0:
                        headers = row_cells
                        table_md_rows.append("| " + " | ".join(headers) + " |")
                        table_md_rows.append("| " + " | ".join(["---"] * len(headers)) + " |")
                    else:
                        table_md_rows.append("| " + " | ".join(row_cells) + " |")

                if table_md_rows:
                    md_table_str = "\n".join(table_md_rows)
                    slide_tables.append(md_table_str)
                    slide_paragraphs.append(md_table_str)

            # Check for text boxes
            elif shape.has_text_frame and shape != slide.shapes.title:
                tf = shape.text_frame
                for p in tf.paragraphs:
                    txt = p.text.strip()
                    if not txt:
                        continue

                    # If slide title was not found via placeholder and this is the first shape
                    # with a prominent position, treat as title
                    if not slide_title and len(slide_paragraphs) == 0:
                        slide_title = txt
                        continue

                    # Check bullet level
                    level = getattr(p, "level", 0) or 0
                    indent = "  " * level
                    bullet_prefix = f"{indent}- " if level > 0 else ""
                    slide_paragraphs.append(f"{bullet_prefix}{txt}")

        # 3. Check for speaker notes
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes_text = ""

        # Use presentation title from first slide if available
        if idx == 1 and slide_title:
            presentation_title = slide_title

        slide_entry = {
            "slide_num": idx,
            "title": slide_title or f"Slide {idx}",
            "content": "\n".join(slide_paragraphs),
            "tables": slide_tables,
            "notes": notes_text,
        }
        slides_data.append(slide_entry)

        # Build clean markdown for this slide
        md_part = [f"## Slide {idx}: {slide_entry['title']}"]
        if slide_entry["content"]:
            md_part.append(slide_entry["content"])
        if notes_text:
            md_part.append(f"> **Speaker Notes:** {notes_text}")

        markdown_sections.append("\n\n".join(md_part))

    full_markdown = f"# {presentation_title}\n\n" + "\n\n---\n\n".join(markdown_sections)

    return {
        "title": presentation_title,
        "structured_text": full_markdown,
        "slides": slides_data,
        "total_slides": len(slides_data),
    }
