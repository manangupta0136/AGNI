from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def xlsx_write(table_data, output_path="output.xlsx"):
    """
    Create a professional Excel file from table data.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = "Inspection Report"

    # Add data
    for row in table_data:
        ws.append(row)

    # Header formatting
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        cell.border = thin_border

    # Format all cells
    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Center important columns
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in [1, 4]:
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="top",
                    wrap_text=True
                )

    # Adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            if cell.value is not None:
                length = len(str(cell.value))
                max_length = max(max_length, length)

        ws.column_dimensions[column_letter].width = min(
            max_length + 3,
            45
        )

    # Freeze header
    ws.freeze_panes = "A2"

    # Make header row slightly taller
    ws.row_dimensions[1].height = 25

    # Add filter
    ws.auto_filter.ref = ws.dimensions

    # Save
    wb.save(output_path)

    return output_path


def xlsx_read(file_path):
    """
    Read an Excel file and return its data.
    """

    wb = load_workbook(file_path)
    ws = wb.active

    return [
        [cell.value for cell in row]
        for row in ws.iter_rows()
    ]